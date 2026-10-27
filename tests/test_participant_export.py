import asyncio
from datetime import date, timedelta
from io import BytesIO
from threading import Event
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4
from zipfile import ZipFile

import pytest
from aiogram.types import User
from openpyxl import load_workbook

from app.events.table import EventTable
from app.miniapp import export_routes
from app.miniapp.state import EXPORT_LOCK, REPOSITORY, USER
from app.services.participant_export import XLSX_CONTENT_TYPE, participants_xlsx
from registrar.models import Student
from tests import test_miniapp
from tests.helpers import complete_student_registration
from tests.test_event_forms import configuration, question
from tests.test_miniapp import headers

miniapp_client = test_miniapp.client


def workbook(content):
    return load_workbook(BytesIO(content))


def test_xlsx_preserves_literal_text_precision_long_answers_and_portable_table():
    questions = [question(label="Name"), question("multiple_choice", label="Track")]
    columns = [
        {"id": "score", "label": "Score", "type": "number"},
        {"id": "large", "label": "Large number", "type": "number"},
        {"id": "day", "label": "Date", "type": "date"},
        {"id": "flag", "label": "Checked", "type": "checkbox"},
        {"id": "formula", "label": "=1+1", "type": "text"},
    ]
    long = "Long answer 🎓\n" * 100
    page = SimpleNamespace(
        event=SimpleNamespace(title="Campus event", questions=questions, table_columns=columns),
        members=[
            {
                "student_id": "00000001",
                "name": "Student\x00 Name",
                "telegram_id": 100,
                "telegram_username": "student_100",
                "email": "student@example.test",
                "phone_number": "+998901234567",
                "education_level": "Level 4",
                "answers": {questions[0]["id"]: long, questions[1]["id"]: ["Option A", "Option B"]},
                "organizer_values": {
                    "score": "8.5",
                    "large": "100000000000000000000000000000000000001",
                    "day": "2026-10-27",
                    "flag": False,
                    "formula": '=HYPERLINK("https://example.test")',
                },
            }
        ],
    )
    content = participants_xlsx(page)
    book = workbook(content)
    sheet = book["Participants"]
    assert sheet.freeze_panes == "C2"
    assert sheet.max_row == 2 and sheet.max_column == 13
    assert [cell.value for cell in sheet[1]][6:] == [
        "Name (2)",
        "Track",
        "Score",
        "Large number",
        "Date",
        "Checked",
        "=1+1",
    ]
    assert sheet["A2"].value == "00000001" and sheet["A2"].data_type == "s"
    assert sheet["B2"].value == "Student Name"
    assert sheet["E2"].value == "+998901234567" and sheet["E2"].data_type == "s"
    assert sheet["G2"].value == long and sheet["G2"].alignment.wrap_text
    assert sheet["H2"].value == "Option A\nOption B"
    assert sheet["I2"].value == 8.5 and sheet["I2"].data_type == "n"
    assert sheet["J2"].value == page.members[0]["organizer_values"]["large"]
    assert sheet["K2"].value.date() == date(2026, 10, 27)
    assert sheet["L2"].value is False
    assert all(cell.data_type != "f" for row in sheet for cell in row)
    assert sheet["M2"].value.startswith("=HYPERLINK")
    table = sheet.tables["Participants"]
    assert table.ref == "A1:M2" and table.autoFilter.ref == table.ref
    assert [column.name for column in table.tableColumns] == [cell.value for cell in sheet[1]]
    with ZipFile(BytesIO(content)) as archive:
        assert not any(
            "vbaProject" in name or "externalLinks" in name for name in archive.namelist()
        )
    book.close()


@pytest.mark.asyncio
async def test_export_contains_every_page_current_cells_and_excludes_expired_answers(
    miniapp_client, repository
):
    client = miniapp_client
    event = await client.events.settings(100, client.event_id)
    q = question(label="Experience")
    event = await client.events.configure(100, event.id, configuration(event, questions=[q]))
    for user_id in range(1000, 1012):
        await repository.touch_user(User(id=user_id, is_bot=False, first_name="Student"))
        await complete_student_registration(
            repository,
            user_id,
            Student(
                str(user_id).zfill(8),
                f"Student {user_id}",
                f"{user_id}@example.test",
                account=f"{user_id}@example.test",
            ),
        )
        await client.events.submit_form(
            user_id, event.id, {q["id"]: f"Response {user_id}"}, event.form_revision
        )
    column = {"id": str(uuid4()), "label": "Score", "type": "number", "options": []}
    table = EventTable(repository.db, clock=repository.clock)
    await table.columns(100, event.id, [column], 1)
    await table.edit(
        100,
        event.id,
        1011,
        {
            "scope": "column",
            "field": column["id"],
            "value": "9",
            "revision": 1,
            "columns_revision": 2,
        },
    )
    path = f"/api/forms/{event.id}/members.xlsx"
    response = await client.get(path + "?page=1", headers=headers())
    assert response.status == 200 and response.content_type == XLSX_CONTENT_TYPE
    assert response.headers["Cache-Control"] == "no-store"
    assert (
        response.headers["Content-Disposition"]
        == f'attachment; filename="participants-{event.id}.xlsx"'
    )
    book = workbook(await response.read())
    sheet = book.active
    assert sheet.max_row == 13
    assert sheet["A2"].value == "00001000" and sheet["A13"].value == "00001011"
    assert sheet["G13"].value == "Response 1011" and sheet["H13"].value == 9
    assert len((await client.events.answer_page(100, event.id)).members) == 10
    book.close()
    repository.test_clock[0] += timedelta(days=31)
    response = await client.get(path, headers=headers())
    book = workbook(await response.read())
    assert book.active.max_row == 13
    assert all(
        book.active.cell(row, column).value is None for row in range(2, 14) for column in (7, 8)
    )
    book.close()


@pytest.mark.asyncio
async def test_empty_export_is_valid_and_download_requires_signed_owner(miniapp_client):
    client = miniapp_client
    assert (
        await client.get(
            f"/api/forms/{client.event_id}/members.xlsx?user_id=100", headers=headers(200)
        )
    ).status == 403
    await client.events.create(200, "Another organizer")
    for resource in ("forms", "events"):
        path = f"/api/{resource}/{client.event_id}/members.xlsx"
        assert (await client.get(path)).status == 401
        assert (await client.get(path + "?user_id=100", headers=headers(200))).status == 400
        response = await client.get(path, headers=headers())
        assert response.status == 200
        book = workbook(await response.read())
        assert book.active.max_row == 1 and book.active["A1"].value == "Student ID"
        book.close()


@pytest.mark.asyncio
async def test_export_runs_off_loop_rejects_overlapping_work_and_releases_slot(
    miniapp_client, monkeypatch
):
    client = miniapp_client
    started, release = Event(), Event()

    def writer(page):
        started.set()
        assert release.wait(timeout=5)
        return participants_xlsx(page)

    monkeypatch.setattr(export_routes, "participants_xlsx", writer)
    path = f"/api/forms/{client.event_id}/members.xlsx"
    task = asyncio.create_task(client.get(path, headers=headers()))
    try:
        assert await asyncio.to_thread(started.wait, 3)
        health = await asyncio.wait_for(client.get("/health"), timeout=1)
        assert health.status == 200
        assert (await client.get(path, headers=headers())).status == 429
    finally:
        release.set()
    response = await task
    assert response.status == 200 and (await response.read()).startswith(b"PK")
    assert not client.server.app[EXPORT_LOCK].locked()


@pytest.mark.asyncio
async def test_export_keeps_slot_until_cancelled_writer_finishes_and_does_not_block_loop(
    monkeypatch,
):
    started, release = asyncio.Event(), asyncio.Event()

    async def writer(*args):
        started.set()
        await release.wait()
        return b"xlsx"

    monkeypatch.setattr(export_routes.asyncio, "to_thread", writer)
    lock = asyncio.Lock()
    request = SimpleNamespace(
        app={
            EXPORT_LOCK: lock,
            REPOSITORY: SimpleNamespace(
                export_members=AsyncMock(
                    return_value=SimpleNamespace(event=SimpleNamespace(id=str(uuid4())))
                )
            ),
        },
        match_info={"event_id": str(uuid4())},
    )

    class Request:
        app = request.app
        match_info = request.match_info

        def __getitem__(self, key):
            assert key is USER
            return 100

    task = asyncio.create_task(export_routes.export_participants(Request()))
    await started.wait()
    assert lock.locked()
    task.cancel()
    await asyncio.sleep(0)
    assert lock.locked()
    release.set()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert not lock.locked()
