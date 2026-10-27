"""Portable XLSX tables with literal user text and no persistent export files."""

import re
from datetime import date
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.event_errors import EventError

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
CONTACT_HEADERS = ["Student ID", "Name", "Telegram", "Email", "Phone", "Level"]


def text(value):
    return INVALID_XML.sub("", str(value))


def headers(labels):
    """Excel table column names must be unique, including case-insensitive duplicates."""
    used, result = set(), []
    for label in labels:
        original = text(label).strip() or "Column"
        candidate, suffix = original, 2
        while candidate.casefold() in used:
            candidate = f"{original} ({suffix})"
            suffix += 1
        used.add(candidate.casefold())
        result.append(candidate)
    return result


def cell_value(value, kind):
    if value is None:
        return None
    if isinstance(value, list):
        return "\n".join(text(item) for item in value)
    if type(value) is bool:
        return value
    if kind == "date":
        return date.fromisoformat(value)
    if kind == "number":
        number = Decimal(value)
        # Excel stores only 15 significant digits. Preserve larger values as text.
        digits = "".join(str(digit) for digit in number.as_tuple().digits).rstrip("0")
        if number.is_finite() and len(digits) <= 15:
            converted = float(number)
            if abs(converted) < float("inf") and (converted != 0 or number == 0):
                return converted
    return text(value)


def participants_xlsx(page):
    """Render an already-authorized, expiry-filtered snapshot outside the event loop."""
    if len(page.members) >= 1_048_576:
        raise EventError("This form exceeds the Excel worksheet row limit")
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Participants"
    workbook.properties.title = text(page.event.title)
    workbook.properties.creator = "WIUT Social Registrar"
    definitions = [*page.event.questions, *page.event.table_columns]
    names = headers([*CONTACT_HEADERS, *(item["label"] for item in definitions)])
    worksheet.append(names)
    worksheet.freeze_panes = "C2"
    worksheet.sheet_view.showGridLines = False
    for index, name in enumerate(names, 1):
        worksheet.column_dimensions[get_column_letter(index)].width = (
            [16, 30, 24, 36, 22, 16][index - 1] if index <= 6 else 45
        )
    for member in page.members:
        values = [
            member["student_id"],
            member["name"],
            "@" + member["telegram_username"]
            if member["telegram_username"]
            else f"tg://user?id={member['telegram_id']}",
            member["email"],
            member["phone_number"],
            member["education_level"],
            *(
                cell_value((member.get("answers") or {}).get(q["id"]), q["type"])
                for q in page.event.questions
            ),
            *(
                cell_value(member["organizer_values"].get(c["id"]), c["type"])
                for c in page.event.table_columns
            ),
        ]
        worksheet.append([text(value) if isinstance(value, str) else value for value in values])
    for row in worksheet:
        for cell in row:
            if isinstance(cell.value, str):
                # Never interpret student responses, labels, IDs or contacts as formulas.
                cell.value = text(cell.value)
                cell.data_type = "s"
                cell.number_format = "@"
            elif isinstance(cell.value, date):
                cell.number_format = "yyyy-mm-dd"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if row[0].row == 1:
            for cell in row:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="217346")
    last = get_column_letter(len(names))
    reference = f"A1:{last}{max(1, worksheet.max_row)}"
    if page.members:
        table = Table(displayName="Participants", ref=reference)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
        worksheet.add_table(table)
    else:
        worksheet.auto_filter.ref = reference
    output = BytesIO()
    try:
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()
        output.close()
